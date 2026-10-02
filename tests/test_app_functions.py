"""Direct offline app contracts complement Streamlit's threaded AppTest coverage."""

import importlib
import io
import json
import sys
from collections.abc import Iterator
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from types import ModuleType, SimpleNamespace
from unittest.mock import MagicMock, Mock, call

import pytest
from docx import Document
from docx.shared import Inches, Pt

from _streamlit_app import app_core


@pytest.fixture
def app_module(monkeypatch: pytest.MonkeyPatch) -> Iterator[ModuleType]:
    """Import the trusted app in the pytest thread with inactive UI and offline services."""
    ui = MagicMock()
    ui.cache_resource = lambda function: function
    ui.button.return_value = False
    ui.toggle.return_value = False
    ui.columns.side_effect = lambda widths: [MagicMock() for _ in widths]
    # Main only asks for membership when no processing button has been pressed.
    ui.session_state = MagicMock(key_textinput="Source")
    config = app_core.load_yaml_config(app_core.repo_path("config.yaml"))
    ui.radio.return_value = config["models"][0]["name"]
    monkeypatch.setitem(sys.modules, "streamlit", ui)
    monkeypatch.setitem(sys.modules, "app_core", app_core)
    monkeypatch.syspath_prepend(str(app_core.APP_DIR))
    monkeypatch.setenv("OPENROUTER_API_KEY", "offline-test-key")
    monkeypatch.setattr("dotenv.load_dotenv", lambda *args, **kwargs: None)
    monkeypatch.setattr(app_core, "start_understandability_loading", lambda: None)
    monkeypatch.setattr(
        app_core, "get_zix", lambda text: {"Source": -1.6, "Result": 1.4}[text]
    )
    monkeypatch.setattr(app_core, "get_cefr", lambda score: {1.4: "B1"}[score])
    name = "_streamlit_app.sprache-vereinfachen"
    monkeypatch.delitem(sys.modules, name, raising=False)
    module = importlib.import_module(name)
    ui.reset_mock()
    ui.session_state.key_textinput = "Source"
    try:
        yield module
    finally:
        sys.modules.pop(name, None)


@pytest.mark.parametrize("easy", [False, True])
@pytest.mark.parametrize("analysis", [False, True])
def test_invoke_model_sends_full_request_and_returns_clean_first_choice(
    app_module: ModuleType, monkeypatch: pytest.MonkeyPatch, easy: bool, analysis: bool
) -> None:
    tag = "leichtesprache" if easy else "einfachesprache"
    complete = Mock(
        return_value=SimpleNamespace(
            choices=[
                SimpleNamespace(
                    message=SimpleNamespace(
                        content=f"  <{tag}># Titel\n**Result**</{tag}>  "
                    )
                ),
                SimpleNamespace(message=SimpleNamespace(content="wrong choice")),
            ]
        )
    )
    monkeypatch.setattr(
        app_module,
        "get_openrouter_client",
        lambda: SimpleNamespace(
            chat=SimpleNamespace(completions=SimpleNamespace(create=complete))
        ),
    )
    monkeypatch.setattr(app_module, "leichte_sprache", easy)
    monkeypatch.setattr(app_module, "condense_text", True)
    monkeypatch.setattr(app_module, "TEMPERATURE", 0.7)
    monkeypatch.setattr(app_module, "MAX_TOKENS", 1234)
    monkeypatch.setattr(
        app_module,
        "MODEL_PARAMETERS",
        {"offline/model": {"extra_body": {"provider": {"only": ["offline"]}}}},
    )
    assert app_module.invoke_model("Source", "offline/model", analysis=analysis) == (
        True,
        "Titel\nResult",
    )
    kwargs = complete.call_args.kwargs
    assert kwargs == {
        "model": "offline/model",
        "extra_body": {"provider": {"only": ["offline"]}},
        "temperature": 0.7,
        "max_tokens": 1234,
        "messages": kwargs["messages"],
    }
    system, user = kwargs["messages"]
    assert system["role"] == "system"
    assert f"<{tag}>" in system["content"]
    assert user["role"] == "user"
    assert user["content"].endswith("Source")
    if analysis:
        assert "Satz für Satz" in user["content"]
    elif easy:
        assert "lass den Rest weg" in user["content"]
    else:
        assert "ALLE Informationen" in user["content"]


@pytest.mark.parametrize(
    "content", [None, "", "untagged", "<einfachesprache> \n </einfachesprache>"]
)
def test_invoke_model_invalid_response_returns_generic_failure(
    app_module: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    content: str | None,
    caplog: pytest.LogCaptureFixture,
) -> None:
    complete = Mock(
        return_value=SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content=content))]
        )
    )
    monkeypatch.setattr(
        app_module,
        "get_openrouter_client",
        lambda: SimpleNamespace(
            chat=SimpleNamespace(completions=SimpleNamespace(create=complete))
        ),
    )
    model = next(iter(app_module.MODEL_PARAMETERS))
    assert app_module.invoke_model("Source", model) == (
        False,
        "Model response could not be created.",
    )
    assert any(
        record.exc_info and record.args == (model,) and record.levelname == "ERROR"
        for record in caplog.records
    )


def test_provider_timeout_is_caught_at_client_boundary(
    app_module: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    complete = Mock(side_effect=TimeoutError("offline timeout"))
    monkeypatch.setattr(
        app_module,
        "get_openrouter_client",
        lambda: SimpleNamespace(
            chat=SimpleNamespace(completions=SimpleNamespace(create=complete))
        ),
    )
    assert app_module.invoke_model(
        "Source", next(iter(app_module.MODEL_PARAMETERS))
    ) == (False, "Model response could not be created.")
    assert any(
        record.exc_info and record.exc_info[0] is TimeoutError
        for record in caplog.records
    )


@pytest.mark.parametrize("easy", [False, True])
def test_response_wrapper_selects_language_tag(
    app_module: ModuleType, monkeypatch: pytest.MonkeyPatch, easy: bool
) -> None:
    monkeypatch.setattr(app_module, "leichte_sprache", easy)
    assert app_module.get_result_from_response(
        "<einfachesprache>Simple</einfachesprache><leichtesprache>Easy</leichtesprache>"
    ) == ("Easy" if easy else "Simple")


def test_one_click_uses_source_for_each_model_and_reports_partial_failure(
    app_module: ModuleType, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(app_module, "MODEL_IDS", {"A": "offline/a", "B": "offline/b"})

    def complete(text: str, model: str, analysis: bool = False) -> tuple[bool, str]:
        assert text == "Source"
        assert analysis is False
        return {"offline/a": (True, "Result"), "offline/b": (False, "private details")}[
            model
        ]

    monkeypatch.setattr(app_module, "invoke_model", complete)
    monkeypatch.setattr(app_module, "get_cefr", lambda score: {1: "B1"}[score])
    assert app_module.get_one_click_results() == (
        True,
        "\n----- Ergebnis von A (Verständlichkeit: 1, Niveau etwa B1) -----\n\nResult"
        "\n\n\n\n----- Fehlgeschlagen: B -----\n\nFür diese Modelle konnte kein Ergebnis erstellt werden.",
    )


def test_sample_text_callback_sets_input(app_module: ModuleType) -> None:
    app_module.enter_sample_text()
    assert app_module.st.session_state.key_textinput == app_module.SAMPLE_TEXT
    assert app_module.st.session_state.key_textinput.startswith(
        "Als Vernehmlassungsverfahren"
    )


def test_project_info_places_image_between_markdown_sections(
    app_module: ModuleType,
) -> None:
    app_module.create_project_info("Before ADD_IMAGE_HERE After")
    ui = app_module.st
    ui.expander.assert_called_once_with("Detaillierte Informationen zum Projekt")
    assert [item for item in ui.mock_calls if item[0] in {"markdown", "image"}] == [
        call.markdown("Before ", unsafe_allow_html=True),
        call.image(str(app_core.app_path("zix_scores.jpg")), width="content"),
        call.markdown(" After", unsafe_allow_html=True),
    ]


@pytest.fixture
def result() -> app_core.ResultState:
    return app_core.ResultState(
        source_text="Source",
        response="Result",
        analysis=False,
        simplification=True,
        one_click=False,
        model_choice="A",
        model_names=("A", "B"),
        time_processed=1.23456,
        score_source=-1.6,
    )


@pytest.mark.parametrize("mode", ["simplification", "analysis", "one_click"])
def test_download_contains_original_output_and_mode_metadata(
    app_module: ModuleType,
    result: app_core.ResultState,
    mode: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result = replace(
        result,
        analysis=mode == "analysis",
        one_click=mode == "one_click",
        simplification=mode == "simplification",
    )
    clock = Mock()
    clock.now.return_value = datetime(2026, 1, 2, 3, 4, 5)
    monkeypatch.setattr(app_module, "datetime", clock)
    app_module.create_download_link(result)
    kwargs = app_module.st.download_button.call_args.kwargs
    assert kwargs["mime"] == app_module.DOWNLOAD_MIME_TYPE
    assert kwargs["file_name"] == (
        app_module.ANALYSIS_FILENAME
        if result.analysis
        else app_module.DEFAULT_OUTPUT_FILENAME
    )
    assert (
        kwargs["label"]
        == {
            "analysis": "Analyse herunterladen",
            "one_click": "Vereinfachte Texte herunterladen",
            "simplification": "Vereinfachten Text herunterladen",
        }[mode]
    )
    document = Document(io.BytesIO(kwargs["data"]))
    heading = {
        "analysis": "Analyse von Sprachmodell A",
        "one_click": "Vereinfachte Texte von Sprachmodellen",
        "simplification": "Vereinfachter Text von Sprachmodell",
    }[mode]
    assert [p.text for p in document.paragraphs] == [
        "Ausgangstext",
        "\nSource",
        heading,
        "Result",
    ]
    for index, paragraph in enumerate(document.paragraphs):
        for run in paragraph.runs:
            assert run.font.name == app_module.FONT_WORDDOC
            assert run.font.size == Pt(
                app_module.FONT_SIZE_HEADING
                if index in {0, 2}
                else app_module.FONT_SIZE_PARAGRAPH
            )
    section = document.sections[0]
    # Word stores page dimensions in twips, so allow that quantization.
    assert abs(section.page_width - Inches(app_module.PAGE_WIDTH_INCHES)) <= 635
    assert abs(section.page_height - Inches(app_module.PAGE_HEIGHT_INCHES)) <= 635
    footer = section.footer.paragraphs[0]
    assert footer.text == (
        f"Erstellt am {clock.now.return_value.strftime(app_module.DATETIME_FORMAT)} mit der Prototyp-App «Einfache Sprache», Amt für Statistik und Daten, Kanton Zürich.\n"
        f"Sprachmodell(e): {'A, B' if result.one_click else 'A'}\nVerarbeitungszeit: 1.2 Sekunden"
    )
    for run in footer.runs:
        assert run.font.name == app_module.FONT_WORDDOC
        assert run.font.size == Pt(app_module.FONT_SIZE_FOOTER)


@pytest.mark.parametrize("mode", ["simplification", "analysis", "one_click"])
def test_render_result_uses_mode_specific_scores_and_offers_download(
    app_module: ModuleType,
    result: app_core.ResultState,
    mode: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result = replace(
        result,
        analysis=mode == "analysis",
        one_click=mode == "one_click",
        simplification=mode == "simplification",
    )
    score = Mock(side_effect=lambda text: {"Result": 1.4}[text])
    monkeypatch.setattr(app_module, "get_zix", score)
    app_module.render_result(result)
    ui = app_module.st
    ui.text_area.assert_called_once_with(
        "Deine Analyse" if result.analysis else "Dein vereinfachter Text",
        height=app_module.TEXT_AREA_HEIGHT,
        value="Result",
    )
    expected = {
        "label": app_module.METRIC_LABEL,
        "help": app_module.METRIC_HELP,
        "value": 1 if result.simplification else -2,
    }
    if result.simplification:
        expected["delta"] = 3
        score.assert_called_once_with("Result")
        assert ":green[Sprachniveau B1]" in ui.markdown.call_args.args[0]
        assert "1 auf einer Skala" in ui.markdown.call_args.args[0]
        assert ui.markdown.call_args.args[0].startswith("Dein vereinfachter Text ist ")
    else:
        score.assert_not_called()
        ui.markdown.assert_not_called()
    ui.metric.assert_called_once_with(**expected)
    ui.caption.assert_called_once_with("Verarbeitet in 1.2 Sekunden.")
    ui.download_button.assert_called_once()


def test_log_event_writes_only_metadata_to_configured_logger(
    app_module: ModuleType, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    destination = tmp_path / "events.log"
    logger = app_core.configure_event_logger(
        {"enabled": True, "filename": str(destination)}
    )
    monkeypatch.setattr(app_module, "EVENT_LOGGER", logger)
    clock = Mock()
    clock.now.return_value = datetime(2026, 1, 2, 3, 4, 5)
    monkeypatch.setattr(app_core, "datetime", clock)
    app_module.log_event(
        "private source",
        "private response",
        True,
        False,
        True,
        True,
        "A",
        1.23456,
        False,
    )
    for handler in logger.handlers:
        handler.flush()
    entry = json.loads(destination.read_text(encoding="utf-8"))
    assert entry["event"] == {
        "timestamp": clock.now.return_value.strftime(app_module.DATETIME_FORMAT),
        "input_chars": 14,
        "response_chars": 16,
        "do_analysis": True,
        "do_simplification": False,
        "do_one_click": True,
        "leichte_sprache": True,
        "model_choice": "A",
        "time_processed_seconds": 1.235,
        "success": False,
    }
    assert entry["logger"] == "simply_simplify_language.events"
    assert "private" not in destination.read_text(encoding="utf-8")


def test_invoke_model_defaults_to_simplification(
    app_module: ModuleType, monkeypatch: pytest.MonkeyPatch
) -> None:
    complete = Mock(
        return_value=SimpleNamespace(
            choices=[
                SimpleNamespace(
                    message=SimpleNamespace(
                        content="<einfachesprache>Result</einfachesprache>"
                    )
                )
            ]
        )
    )
    monkeypatch.setattr(
        app_module,
        "get_openrouter_client",
        lambda: SimpleNamespace(
            chat=SimpleNamespace(completions=SimpleNamespace(create=complete))
        ),
    )
    assert app_module.invoke_model(
        "Source", next(iter(app_module.MODEL_PARAMETERS))
    ) == (True, "Result")
    prompt = complete.call_args.kwargs["messages"][1]["content"]
    assert "Satz für Satz" not in prompt
    assert "ALLE Informationen" in prompt


def test_one_click_starts_every_model_before_waiting_for_results(
    app_module: ModuleType, monkeypatch: pytest.MonkeyPatch
) -> None:
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event, Lock

    # A default executor caps its pool at 32. A larger configured model set must
    # still fan out in parallel rather than queue requests behind slow providers.
    models = {f"Model {index}": f"offline/{index}" for index in range(33)}
    monkeypatch.setattr(app_module, "MODEL_IDS", models)
    monkeypatch.setattr(app_module, "get_cefr", lambda score: {1: "B1"}[score])
    started: set[str] = set()
    lock, all_started, release = Lock(), Event(), Event()

    def complete(text: str, model: str) -> tuple[bool, str]:
        assert text == "Source"
        with lock:
            started.add(model)
            if len(started) == len(models):
                all_started.set()
        if not release.wait(5):
            raise TimeoutError("Test did not release model calls")
        return True, "Result"

    monkeypatch.setattr(app_module, "invoke_model", complete)
    with ThreadPoolExecutor(max_workers=1) as caller:
        result = caller.submit(app_module.get_one_click_results)
        try:
            assert all_started.wait(2), (
                "A configured model was queued behind another request"
            )
        finally:
            release.set()
        success, output = result.result(timeout=5)
    assert success is True
    assert output.count("Ergebnis von Model ") == len(models)
